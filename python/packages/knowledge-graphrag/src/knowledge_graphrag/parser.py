import re
import io
from typing import List

from .types import DocumentMime, ParsedDocument, ParsedImageRef, ParsedSection

# 支持的文档 MIME 白名单
SUPPORTED_MIMES: List[str] = [
    "text/plain",
    "text/markdown",
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
]


def is_supported_mime(mime: str) -> bool:
    """检查 MIME 类型是否受支持"""
    return mime in SUPPORTED_MIMES


def _normalize_image_path(raw_path: str) -> str:
    """把 markdown 图片引用中的 path 规范化：去除前导 ./ 与 /、压缩多余斜杠"""
    path = raw_path.replace("\\", "/")
    if path.startswith("./"):
        path = path[2:]
    path = re.sub(r"^/+", "", path)
    path = re.sub(r"/{2,}", "/", path)
    return path.strip()


def _is_external_image(path: str) -> bool:
    """仅扫描 http/https/data: 这类非本地图片引用，命中即返回 true"""
    return bool(re.match(r"^(https?:|data:)", path, re.IGNORECASE))


def parse_text_document(data: bytes, mime: DocumentMime) -> ParsedDocument:
    """解析纯文本 / Markdown 文档"""
    if mime not in ("text/plain", "text/markdown"):
        raise ValueError("Unsupported MIME type")
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        raise ValueError("Invalid UTF-8 document")
    if "\0" in text:
        raise ValueError("binary content is not supported")

    sections: List[ParsedSection] = []
    image_refs: List[ParsedImageRef] = []

    if mime == "text/markdown":
        # 标题按行扫描；图片按全文扫描（图片可与正文同行）
        for line in re.split(r"\r?\n", text):
            match = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
            if match:
                sections.append(ParsedSection(level=len(match.group(1)), title=match.group(2)))

        # ![alt](path "title") —— 允许可选的 title 部分；只关心 path 与 alt
        image_pattern = re.compile(r'!\[((?:[^\]\\]|\\.)*)\]\(\s*([^)\s]+)(?:\s+"[^"]*")?\s*\)')
        for match in image_pattern.finditer(text):
            alt = match.group(1) or ""
            raw_path = match.group(2) or ""
            if not raw_path or _is_external_image(raw_path):
                continue
            image_refs.append(
                ParsedImageRef(path=_normalize_image_path(raw_path), alt=alt, offset=match.start())
            )

    return ParsedDocument(text=text, mime=mime, sections=sections, image_refs=image_refs)


def parse_pdf_document(data: bytes) -> ParsedDocument:
    """从 PDF 字节中提取纯文本（逐页拼接）"""
    try:
        import pymupdf
    except ImportError:
        raise ImportError("pymupdf is required for PDF parsing")

    doc = pymupdf.open(stream=data, filetype="pdf")
    parts: List[str] = []
    for page in doc:
        parts.append(page.get_text())
    doc.close()
    return ParsedDocument(
        text="\n".join(parts),
        mime="application/pdf",
        sections=[],
        image_refs=[],
    )


def parse_docx_document(data: bytes) -> ParsedDocument:
    """从 DOCX 字节中提取纯文本"""
    try:
        from docx import Document as DocxDocument
    except ImportError:
        raise ImportError("python-docx is required for DOCX parsing")

    doc = DocxDocument(io.BytesIO(data))
    parts: List[str] = []
    for para in doc.paragraphs:
        parts.append(para.text)
    return ParsedDocument(
        text="\n".join(parts),
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        sections=[],
        image_refs=[],
    )


def parse_xlsx_document(data: bytes) -> ParsedDocument:
    """从 XLSX 字节中提取所有工作表的文本，每个工作表输出：表名 + CSV 文本"""
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required for XLSX parsing")

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: List[str] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        rows: List[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell) for cell in row]
            rows.append(",".join(cells))
        csv_text = "\n".join(rows)
        parts.append(f"# {sheet_name}\n{csv_text}")
    workbook.close()
    return ParsedDocument(
        text="\n\n".join(parts),
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        sections=[],
        image_refs=[],
    )


def parse_document(data: bytes, mime: DocumentMime) -> ParsedDocument:
    """按 MIME 类型分流解析"""
    if mime in ("text/plain", "text/markdown"):
        return parse_text_document(data, mime)
    elif mime == "application/pdf":
        return parse_pdf_document(data)
    elif mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return parse_docx_document(data)
    elif mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return parse_xlsx_document(data)
    else:
        raise ValueError(f"Unsupported MIME type: {mime}")
